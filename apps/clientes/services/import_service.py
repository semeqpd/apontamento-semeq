import csv
import openpyxl
import unicodedata
import re
from io import BytesIO


# Character replacements for special chars (shared)
_REPLACEMENTS = {
    '\u2013': '-',  # en-dash
    '\u2014': '-',  # em-dash
    '\u2018': "'",  # left single quote
    '\u2019': "'",  # right single quote
    '\u201c': '"',  # left double quote
    '\u201d': '"',  # right double quote
    '\u2026': '...', # ellipsis
    '\u00a0': ' ',  # non-breaking space
    '\u00b7': '-',  # middle dot
    '\u2022': '-',  # bullet
    '\u25aa': '-',  # black small square
    '\u25cf': '-',  # black circle
    '\u0096': '-',  # START OF GUARDED AREA (control char)
    '\u0097': '-',  # END OF GUARDED AREA (control char)
}


def _sanitize_text(text: str) -> str:
    """Normalize unicode and replace special characters."""
    if not text:
        return ''
    text = unicodedata.normalize('NFKC', text)
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    # Remove any remaining control characters (except newline/tab)
    text = re.sub(r'[\x00-\x08\x0b-\x1f\x7f-\x9f]', '', text)
    return text.strip()


def parse_file(arquivo, ext: str):
    """Parse CSV or XLSX file and return headers and rows. Auto-detects CSV delimiter."""
    if ext == 'csv':
        # Try multiple encodings with fallback
        raw_content = arquivo.read()
        encodings = ['utf-8-sig', 'utf-8', 'iso-8859-1', 'latin-1']
        content = None
        for enc in encodings:
            try:
                content = raw_content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if content is None:
            content = raw_content.decode('utf-8', errors='replace')
        
        # Normalize content to fix special characters
        content = unicodedata.normalize('NFKC', content)
        
        # Auto-detect delimiter: comma or semicolon
        sample = content[:1024]
        delimiter = ';' if sample.count(';') > sample.count(',') else ','
        reader = csv.reader(content.splitlines(), delimiter=delimiter)
        headers = next(reader, [])
        rows = list(reader)
    elif ext in ('xlsx', 'xls'):
        wb = openpyxl.load_workbook(BytesIO(arquivo.read()))
        ws = wb.active
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows.append(list(row))
    else:
        raise ValueError(f'Formato não suportado: {ext}')
    return headers, rows


def normalize_row(row, headers):
    """Normalize row data to dict with header keys (case-insensitive)."""
    # Create case-insensitive mapping
    header_map = {h.strip().lower(): h for h in headers if h}
    result = {}
    for key, orig_header in header_map.items():
        idx = headers.index(orig_header)
        val = (row[idx] if idx < len(row) else '') if row else ''
        # Normalize value
        if isinstance(val, str):
            val = _sanitize_text(val)
        result[key] = val
    return result


def fix_encoding(text: str) -> str:
    """Clean text - normalize and replace special chars."""
    return _sanitize_text(text)